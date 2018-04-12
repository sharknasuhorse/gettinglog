import telnetlib
import sys
import time
import __main__
import openpyxl as px

def login ( hostname , address, password ):
    c = telnetlib.Telnet(address,timeout=3)
    print("logging")
    c.read_until(b"Password:")
    c.write(b"\n")
    c.read_until(b"Password:")
    c.write(bytes(password,'utf-8') + b"\n")
    prompt = hostname[:-3] + ">"
    c.read_until(prompt.encode('utf-8'))
    c.write(b"enable\n")
    c.read_until(b"Password:")
    c.write(bytes(password,'utf-8') + b"\n")
    print("fin login")
    __main__.c = c

def length0 (hostname):
    c = __main__.c
    prompt = hostname + "#"
    c.read_until(prompt.encode('utf-8'))
    c.write(b"ter len 0\n")    


def execute_command(hostname,command, teamnum):
    c = __main__.c
    prompt = hostname + "#"
    c.read_until(prompt.encode('utf-8'))
    command = command + "\n"
    c.write(command.encode('utf-8'))
    result = c.read_until(prompt.encode('utf-8'))
    logname = "./" + str(int(teamnum/6+1)) + "-"  + hostname + ".cfg"
    wb = open(logname ,'wb')
    wb.write(result)
    wb.close()
    c.write(b"\n")


def close_connection():
    c = __main__.c
    c.write(b"exit\n")


def get_excel_hostaddr():
    wb = px.load_workbook('testbook.xlsx')
    ws = wb.active
    i = 1
    dct = {}
    for i in range(1,90):
        #print(ws['A' + str(i)].value)
        dct[ws['A' + str(i)].value + "-" +"{0:02d}".format(i)] = ws['B' + str(i)].value
    __main__.dct = dct
    __main__.i = i


def execute_command2(hostname,command):
    c = __main__.c
    prompt = hostname + "#"
    c.read_until(prompt.encode('utf-8'))
    command = command + "\n"
    c.write(command.encode('utf-8'))
    result = c.read_until(prompt.encode('utf-8'))
    logname = "./" + hostname + "-tech-support.log"
    wb = open(logname ,'wb')
    wb.write(result)
    wb.close()
    c.write(b"\n")
passlist = [ "Cnj4BKzr","Dan3ckyA","q0bC50xE","ZWKKEFzV","JxYmgogc","42AUVvLA"]


if __name__ == '__main__':
    get_excel_hostaddr()
    key_lst = list(dct.keys())
    j = 0
    while j < i - 1 :
        try:
            print(key_lst[j][:-3])
            login( key_lst[j], dct[key_lst[j]],passlist[j % 6] )
        except:
            print(dct[key_lst[j]])
            continue
        print("len0")
        length0(key_lst[j][:-3])
        execute_command( (key_lst[j][:-3]),"show run",j)
        close_connection()
        j += 1
        print("next while")
    exit()





